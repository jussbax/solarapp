from __future__ import annotations

from typing import Optional

import logging
import re
import unicodedata
from dataclasses import asdict
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Response, UploadFile
from fastapi.responses import FileResponse
from sqlmodel import Session, select

from ..auth import require_account, require_user
from ..compute import ComputeError, compute_results
from ..config import Settings, get_settings
from ..core.dataset import NasaReference, PvgisDataset
from ..db import get_session
from ..models import Assessment, User, utcnow
from ..pricing.config import settings_version
from ..pricing.datasheets import datasheet_sources
from ..pricing.job import PricingContext
from ..pricing.store import load_catalog, load_config
from ..reports.card import build_client_card
from ..reports.customer_pdf import build_customer_pdf
from ..reports.plans_pdf import build_plans_pdf
from ..reports.program_pdf import build_program_pdf
from ..reports import vicinity
from ..reports.quotation_pdf import build_quotation_pdf, customer_battery_kwh
from ..schemas import JOB_STAGES, AssessmentDoc, AssessmentOut, AssessmentSummary, RevisionEntry, RevisionIn, VicinityMap
from .appliances import remember_appliances
from .deps import get_nasa, get_pvgis
from .settings_routes import company_settings

log = logging.getLogger("solarapp.audit")


def _download_name(prefix: str, a: Assessment, ext: str, inline: bool = False) -> dict:
    """A Content-Disposition that survives quotes and non-Latin names (customer names come from the public form)."""
    raw = (a.customer_name or f"assessment-{a.id}").strip()
    ascii_name = re.sub(r"[^A-Za-z0-9_-]+", "_", unicodedata.normalize("NFKD", raw).encode("ascii", "ignore").decode()).strip("_")[:60] or f"assessment-{a.id}"
    kind = "inline" if inline else "attachment"
    return {"Content-Disposition": f"{kind}; filename=\"{prefix}-{ascii_name}.{ext}\"; filename*=UTF-8''{quote(f'{prefix}-{raw}.{ext}')}"}


router = APIRouter(prefix="/api/assessments", tags=["assessments"], dependencies=[Depends(require_user)])


def _get(session: Session, assessment_id: int) -> Assessment:
    a = session.get(Assessment, assessment_id)
    if a is None:
        raise HTTPException(status_code=404, detail="Assessment not found")
    return a


def _settings_changed(results: Optional[dict], current_version: str) -> bool:
    """The price was calculated under other pricing settings than today's. A record priced before versions were
    stored carries none and is not flagged; its next Calculate stores one."""
    pricing = (results or {}).get("pricing") or {}
    stored = pricing.get("settings_version")
    return bool(pricing.get("available") and stored and stored != current_version)


def project_status(a: Assessment) -> str:
    """The engineering status (round 4), read from the record's facts and never typed: proposal issued (the proposal PDF
    was generated and the design not reopened since), designed (results calculated and not stale), surveyed (at least
    one reading set saved), else draft. The job stage the document keeps is the CRM's and the PM module's, not this."""
    if a.proposal_issued_at is not None:
        return "proposal_issued"
    if a.results and not a.results_stale:
        return "designed"
    if (a.doc or {}).get("reading_sets"):
        return "surveyed"
    return "draft"


def _revisions(a: Assessment) -> list[RevisionEntry]:
    """The revision log as stored (round 13); an older record without the column reads as an empty log."""
    return [RevisionEntry.model_validate(r) for r in (a.revisions or [])]


def _out(a: Assessment, settings_changed: bool = False) -> AssessmentOut:
    return AssessmentOut(
        id=a.id, created_at=a.created_at, updated_at=a.updated_at,
        doc=AssessmentDoc.model_validate(a.doc), results=a.results, results_stale=a.results_stale,
        pricing_settings_changed=settings_changed, status=project_status(a), proposal_issued_at=a.proposal_issued_at,
        plans_issued_at=a.plans_issued_at, revisions=_revisions(a),
        vicinity_map=VicinityMap.model_validate(a.vicinity_map) if a.vicinity_map else None,
    )


def _out_live(session: Session, a: Assessment) -> AssessmentOut:
    """The record with the settings-changed flag judged against the pricing settings as they are now."""
    return _out(a, _settings_changed(a.results, settings_version(load_config(session))))


@router.get("", response_model=list[AssessmentSummary])
def list_assessments(session: Session = Depends(get_session)) -> list[AssessmentSummary]:
    """The project list: engineering facts only. The lead's contact, source and what they saw live on the booking (/api/leads, the CRM's)."""
    rows = session.exec(select(Assessment).order_by(Assessment.updated_at.desc())).all()
    current_version = settings_version(load_config(session))
    out = []
    for a in rows:
        results = a.results or {}
        prod = results.get("production") or {}
        sizing = results.get("sizing") or {}
        if sizing:  # the sized system rather than the roof maximum
            prod = {"system_kwp": sizing.get("kwp"), "annual_kwh": sizing.get("annual_production_kwh"), "total_panels": sizing.get("panels")}
        pricing = results.get("pricing") or {}
        doc = a.doc or {}
        battery = None
        if sizing and (sizing.get("kind") or "") != "net_metering":
            try:
                battery = customer_battery_kwh(pricing if pricing.get("available") else {}, sizing) or None
            except (KeyError, TypeError, ValueError):
                battery = None
        stage = (doc.get("program") or {}).get("stage", "assessed")
        out.append(AssessmentSummary(
            id=a.id, created_at=a.created_at, updated_at=a.updated_at, customer_name=a.customer_name,
            address=a.address, has_results=a.results is not None, results_stale=a.results_stale,
            pricing_settings_changed=_settings_changed(results, current_version),
            stage=stage if stage in JOB_STAGES else "assessed", status=project_status(a), proposal_issued_at=a.proposal_issued_at,
            contract_php=(pricing.get("totals") or {}).get("contract_rounded") if pricing.get("available") else None,
            system_kwp=prod.get("system_kwp"), annual_kwh=prod.get("annual_kwh"), panel_count=prod.get("total_panels"),
            kind=((doc.get("audit") or {}).get("system") or {}).get("kind"),
            face_count=len(doc.get("faces") or []), battery_kwh=battery, computed_at=results.get("computed_at") if a.results else None,
            lead_id=doc.get("lead_id"),
        ))
    return out


@router.post("", response_model=AssessmentOut, status_code=201)
def create_assessment(doc: AssessmentDoc, session: Session = Depends(get_session)) -> AssessmentOut:
    a = Assessment(customer_name=doc.customer_name, address=doc.address, doc=doc.model_dump(mode="json"))
    session.add(a)
    session.commit()
    session.refresh(a)
    return _out(a)


@router.get("/{assessment_id}", response_model=AssessmentOut)
def get_assessment(assessment_id: int, session: Session = Depends(get_session)) -> AssessmentOut:
    return _out_live(session, _get(session, assessment_id))


@router.put("/{assessment_id}", response_model=AssessmentOut)
def update_assessment(assessment_id: int, doc: AssessmentDoc, session: Session = Depends(get_session)) -> AssessmentOut:
    a = _get(session, assessment_id)
    new_doc = doc.model_dump(mode="json")
    stored = AssessmentDoc.model_validate(a.doc or {}).model_dump(mode="json")   # as the browser received it (old panel fields migrated)
    if new_doc != stored:
        # the card's next-step line is printed, never computed: changing it leaves the results fresh
        affects_results = {k: v for k, v in new_doc.items() if k != "card_next_step"} != {k: v for k, v in stored.items() if k != "card_next_step"}
        a.doc = new_doc
        a.customer_name, a.address = doc.customer_name, doc.address
        if affects_results:
            a.results_stale = a.results is not None
        a.updated_at = utcnow()
        session.add(a)
        session.commit()
        session.refresh(a)
        remember_appliances(session, doc)
    return _out_live(session, a)


@router.delete("/{assessment_id}", status_code=204)
def delete_assessment(assessment_id: int, session: Session = Depends(get_session), settings: Settings = Depends(get_settings)) -> Response:
    a = _get(session, assessment_id)
    session.delete(a)
    session.commit()
    vicinity.remove_project_files(settings.data_dir, assessment_id)   # the project's folder (the maps, the upload) goes with the record
    log.info("assessment deleted id=%s", assessment_id)
    return Response(status_code=204)


REPRICE_DETAIL = "Pricing settings changed since the proposal was issued at this price (was ₱{was:,.0f}). Re-pricing a job whose proposal is issued needs your confirmation."


def _needs_reprice_confirmation(a: Assessment, current_version: str) -> Optional[float]:
    """Once the proposal is issued (round 4; the job stage "quoted" before) a price calculated under other pricing
    settings stays until the re-price is confirmed: returns the proposed contract when confirmation is needed, else None."""
    if not _settings_changed(a.results, current_version):
        return None
    if a.proposal_issued_at is None:
        return None
    return float(((a.results or {}).get("pricing") or {}).get("totals", {}).get("contract_rounded") or 0)


@router.post("/{assessment_id}/compute", response_model=AssessmentOut)
def compute_assessment(
    assessment_id: int,
    doc: Optional[AssessmentDoc] = None,
    confirm_reprice: bool = False,
    session: Session = Depends(get_session),
    pvgis: PvgisDataset = Depends(get_pvgis),
    nasa: NasaReference = Depends(get_nasa),
    settings: Settings = Depends(get_settings),
) -> AssessmentOut:
    """Save the document (if sent) and compute results. The pricing settings' version is stored with a priced result;
    when the settings have moved since, a job whose proposal is issued is re-priced only with `confirm_reprice`."""
    a = _get(session, assessment_id)
    if doc is not None:
        a.doc = doc.model_dump(mode="json")
        a.customer_name, a.address = doc.customer_name, doc.address
    parsed = AssessmentDoc.model_validate(a.doc)
    cfg = load_config(session)
    current_version = settings_version(cfg)
    was = _needs_reprice_confirmation(a, current_version)
    if was is not None and not confirm_reprice:
        raise HTTPException(status_code=409, detail=REPRICE_DETAIL.format(was=was))
    ctx = PricingContext(load_catalog(session), cfg, company_settings(session, settings))
    try:
        a.results = compute_results(parsed, pvgis, nasa, ctx)
    except ComputeError as e:
        raise HTTPException(status_code=422, detail=str(e))
    if (a.results.get("pricing") or {}).get("available"):
        a.results["pricing"]["settings_version"] = current_version
    if was is not None:
        log.info("assessment re-priced under new pricing settings id=%s status=%s was=%s", assessment_id, project_status(a), was)
    # the record keeps the document as calculated: the old panel fields are gone and the panels typed by hand were reported once
    parsed.dropped_panels = []
    a.doc = parsed.model_dump(mode="json")
    a.results_stale = False
    a.updated_at = utcnow()
    session.add(a)
    session.commit()
    session.refresh(a)
    remember_appliances(session, parsed)
    return _out(a)


STALE_DETAIL = "Inputs changed since the last calculation. Calculate again first."


def _fresh_results(a: Assessment) -> dict:
    """The one stale rule for every document: results must exist and match the saved inputs."""
    if not a.results:
        raise HTTPException(status_code=409, detail="Calculate first.")
    if a.results_stale:
        raise HTTPException(status_code=409, detail=STALE_DETAIL)
    return a.results


def _customer_results(a: Assessment) -> dict:
    """The customer documents (proposal, roof check, card) add the round-3 design rule to the stale rule: a failed
    circuit coordination (pricing.design_blocked, from the hard warnings) refuses them with the reason, so a wrong
    design never prints as if it were right."""
    results = _fresh_results(a)
    blocked = (results.get("pricing") or {}).get("design_blocked") or []
    if blocked:
        raise HTTPException(status_code=409, detail=f"design_blocked: the design has a hard warning ({', '.join(blocked)}). Fix it under System design, then Calculate again.")
    return results


@router.get("/{assessment_id}/report.pdf")
def customer_report(
    assessment_id: int,
    session: Session = Depends(get_session),
    settings: Settings = Depends(get_settings),
) -> Response:
    a = _get(session, assessment_id)
    results = _customer_results(a)
    if (results.get("dataset") or {}).get("synthetic"):
        raise HTTPException(status_code=409, detail="Customer documents are disabled while test weather data is in use.")
    company = company_settings(session, settings)
    pdf = build_customer_pdf(AssessmentDoc.model_validate(a.doc), results, company)
    return Response(pdf, media_type="application/pdf", headers=_download_name("roof-check", a, "pdf"))


@router.get("/{assessment_id}/quotation.pdf")
def customer_quotation(
    assessment_id: int,
    session: Session = Depends(get_session),
    settings: Settings = Depends(get_settings),
) -> Response:
    a = _get(session, assessment_id)
    results = _customer_results(a)
    if (results.get("dataset") or {}).get("synthetic"):
        raise HTTPException(status_code=409, detail="Customer documents are disabled while test weather data is in use.")
    if not (results.get("pricing") or {}).get("available"):
        raise HTTPException(status_code=409, detail="Calculate first. Pricing needs the panel linked to the materials list.")
    company = company_settings(session, settings)
    pdf = build_quotation_pdf(AssessmentDoc.model_validate(a.doc), results, company, proposal_no=f"P-{a.created_at.year}-{a.id:04d}",
                              outage_hours=load_config(session).program.installation_outage_hours)
    if a.proposal_issued_at is None:
        # the proposal exists for the record from here (round 4): the status reads "proposal issued" with this date and the
        # price locks; "Reopen design" (POST .../reopen) clears it. A later download keeps the first date.
        a.proposal_issued_at = utcnow()
        session.add(a)
        session.commit()
        log.info("proposal issued id=%s", assessment_id)
    return Response(pdf, media_type="application/pdf", headers=_download_name("proposal", a, "pdf"))


@router.post("/{assessment_id}/reopen", response_model=AssessmentOut)
def reopen_design(assessment_id: int, session: Session = Depends(get_session)) -> AssessmentOut:
    """"Reopen design": the proposal issued for this record is no longer the standing one, so the status falls back to
    the facts (designed, surveyed or draft) and Calculate re-prices freely again. The document and the results stay."""
    a = _get(session, assessment_id)
    if a.proposal_issued_at is not None:
        log.info("design reopened id=%s proposal_issued_at=%s", assessment_id, a.proposal_issued_at.isoformat())
        a.proposal_issued_at = None
        a.updated_at = utcnow()
        session.add(a)
        session.commit()
        session.refresh(a)
    return _out_live(session, a)


@router.get("/{assessment_id}/program.pdf")
def program_of_works(
    assessment_id: int,
    session: Session = Depends(get_session),
    settings: Settings = Depends(get_settings),
) -> Response:
    a = _get(session, assessment_id)
    results = _fresh_results(a)
    if not (results.get("program") or {}).get("available"):
        raise HTTPException(status_code=409, detail="Calculate first. Pricing needs the panel linked to the materials list.")
    company = company_settings(session, settings)
    pdf = build_program_pdf(AssessmentDoc.model_validate(a.doc), results, company)
    return Response(pdf, media_type="application/pdf", headers=_download_name("program-of-works", a, "pdf"))


@router.get("/{assessment_id}/plans.pdf")
def plans_for_the_pee(
    assessment_id: int,
    session: Session = Depends(get_session),
    settings: Settings = Depends(get_settings),
) -> Response:
    """The plans for the PEE (round 4): the A3 drawing set the signing engineer seals, built from the geometry, the
    BOM and the settings of the current calculation. Internal, so it prints on test weather like the program of works;
    refused on stale or design-blocked results the way the proposal is, and without pricing (no BOM, no circuits)."""
    a = _get(session, assessment_id)
    results = _customer_results(a)
    if not (results.get("pricing") or {}).get("available"):
        raise HTTPException(status_code=409, detail="Calculate first. Pricing needs the panel linked to the materials list.")
    company = company_settings(session, settings)
    catalog = load_catalog(session, include_inactive=True)
    items = {code: asdict(item) for code, item in catalog.items.items()}
    if a.plans_issued_at is None:
        # the set exists for the record from here (round 13, brief 6.3): revision 0, "first issue", dated now; a later
        # download keeps the first date, and "Issue a revision" appends the next number. Nothing else about the record moves.
        a.plans_issued_at = utcnow()
        session.add(a)
        session.commit()
        session.refresh(a)
        log.info("plans issued id=%s", assessment_id)
    # the vicinity map on record (round 13, item 4): never fetched here, so the build never waits on the tile server
    pdf = build_plans_pdf(AssessmentDoc.model_validate(a.doc), results, company, items=items, config=load_config(session).model_dump(mode="json"),
                          project_no=f"P-{a.created_at.year}-{a.id:04d}", datasheets=datasheet_sources(session),
                          plans_issued_at=a.plans_issued_at.isoformat(), revisions=[r.model_dump() for r in _revisions(a)],
                          vicinity=a.vicinity_map, project_dir=vicinity.project_dir(settings.data_dir, a.id))
    return Response(pdf, media_type="application/pdf", headers=_download_name("plans", a, "pdf"))


@router.post("/{assessment_id}/revisions", response_model=AssessmentOut)
def issue_revision(assessment_id: int, body: RevisionIn, account: User = Depends(require_account), session: Session = Depends(get_session)) -> AssessmentOut:
    """"Issue a revision" (round 13, brief 6.3): appends the next number to the plan set's revision log with the note and
    the signed-in person's name. The log is append-only; revision 0 is the first issue, so the plans must have been
    generated once before a revision can be issued. The next plans PDF prints the new number on every sheet."""
    a = _get(session, assessment_id)
    if a.plans_issued_at is None:
        raise HTTPException(status_code=409, detail="Generate the plans for the PEE first: the first issue is revision 0, and a revision follows it.")
    note = body.note.strip()
    if not note:
        raise HTTPException(status_code=422, detail="A revision needs a note that says what changed.")
    entries = list(a.revisions or [])
    entry = RevisionEntry(no=len(entries) + 1, date=utcnow().isoformat(), note=note, by=(account.display_name or account.username).strip())
    a.revisions = entries + [entry.model_dump()]   # a new list, so SQLAlchemy sees the JSON column change
    a.updated_at = utcnow()
    session.add(a)
    session.commit()
    session.refresh(a)
    log.info("plans revision issued id=%s no=%s by=%s", assessment_id, entry.no, account.username)
    return _out_live(session, a)


# ---- the vicinity map (round 13, item 4; docs/audits/round-13/engineer-brief.md 4.1): the composed map tiles and the office's upload

def _store_vicinity(session: Session, a: Assessment, state: dict) -> AssessmentOut:
    a.vicinity_map = dict(state)   # a new dict, so SQLAlchemy sees the JSON column change
    session.add(a)
    session.commit()
    session.refresh(a)
    return _out_live(session, a)


@router.post("/{assessment_id}/vicinity-map/fetch", response_model=AssessmentOut)
def prepare_vicinity_map(assessment_id: int, force: bool = Query(False, description="remake the mosaics even when they were made for this pin"),
                         session: Session = Depends(get_session), settings: Settings = Depends(get_settings)) -> AssessmentOut:
    """"Prepare the map": composes the vicinity map from map tiles for the record's pin under the tile usage policy
    (one request at a time, the company's contact in the User-Agent, the 30-day tile cache) and keeps it with the project;
    nothing is fetched when the mosaics on record were made for this pin (unless `force`). A failure (no outside
    access, a timeout, a 429) is kept with its reason, and the sheet prints the pin and that reason instead of a map.
    The plans build itself never fetches, so it never waits on the network."""
    a = _get(session, assessment_id)
    doc = AssessmentDoc.model_validate(a.doc)
    if doc.lat is None or doc.lon is None:
        raise HTTPException(status_code=409, detail="Set the map pin first: the vicinity map is composed around it.")
    if vicinity.FETCH_LOCK.locked():
        raise HTTPException(status_code=409, detail="A vicinity map is being prepared already; try again in a moment.")
    agent = vicinity.user_agent(company_settings(session, settings), settings.website_base)
    with vicinity.make_client() as client:   # the environment's proxy settings apply, as to every outbound call
        state = vicinity.prepare(doc.lat, doc.lon, data_dir=settings.data_dir, assessment_id=a.id, url_template=settings.map_tiles_url,
                                 attribution=settings.map_tiles_attribution, agent=agent, state=a.vicinity_map, client=client, force=force)
    return _store_vicinity(session, a, state)


@router.post("/{assessment_id}/vicinity-map", response_model=AssessmentOut)
async def upload_vicinity_map(assessment_id: int, file: UploadFile = File(...), note: str = Form(""),
                              session: Session = Depends(get_session), settings: Settings = Depends(get_settings)) -> AssessmentOut:
    """The office's screen grab of the vicinity map (the override: it is printed in place of the fetched map whenever
    it is present, with the attribution typed as `note`): PNG or JPEG, 8 MB at most, re-encoded without EXIF and the
    longer side capped at 2,400 px, kept with the project."""
    a = _get(session, assessment_id)
    if file.size is not None and file.size > vicinity.UPLOAD_MAX_BYTES:
        raise HTTPException(status_code=413, detail="The image is larger than 8 MB.")
    data = await file.read(vicinity.UPLOAD_MAX_BYTES + 1)
    if len(data) > vicinity.UPLOAD_MAX_BYTES:
        raise HTTPException(status_code=413, detail="The image is larger than 8 MB.")
    try:
        state = vicinity.store_upload(data, note, vicinity.project_dir(settings.data_dir, a.id), state=a.vicinity_map)
    except vicinity.UploadError as e:
        raise HTTPException(status_code=422, detail=str(e))
    log.info("vicinity map uploaded id=%s bytes=%s", assessment_id, len(data))
    return _store_vicinity(session, a, state)


@router.delete("/{assessment_id}/vicinity-map", response_model=AssessmentOut)
def remove_vicinity_upload(assessment_id: int, session: Session = Depends(get_session), settings: Settings = Depends(get_settings)) -> AssessmentOut:
    """Removes the office's upload; the fetched map, when on record, prints again."""
    a = _get(session, assessment_id)
    state = vicinity.remove_upload(a.vicinity_map, vicinity.project_dir(settings.data_dir, a.id))
    log.info("vicinity map upload removed id=%s", assessment_id)
    return _store_vicinity(session, a, state)


@router.get("/{assessment_id}/vicinity-map.png")
def vicinity_map_png(assessment_id: int, which: str = Query("", pattern="^(|upload|z16|z12)$"),
                     session: Session = Depends(get_session), settings: Settings = Depends(get_settings)) -> Response:
    """The map the sheet prints (the upload when present, else the main mosaic), for the Site plan card's preview;
    `which` asks for one file outright. 404 when nothing is on record."""
    a = _get(session, assessment_id)
    path = vicinity.current_file(a.vicinity_map, vicinity.project_dir(settings.data_dir, a.id), which)
    if path is None:
        raise HTTPException(status_code=404, detail="No vicinity map on record for this project.")
    return FileResponse(path, media_type="image/png", headers={"Cache-Control": "no-store"})


@router.get("/{assessment_id}/card.png")
def client_card(
    assessment_id: int,
    next_step: Optional[str] = None,
    session: Session = Depends(get_session),
    settings: Settings = Depends(get_settings),
) -> Response:
    """Phone-sized image of the roof check result for the customer. The next step comes from the query
    string when given, else from the record's saved card_next_step."""
    a = _get(session, assessment_id)
    results = _customer_results(a)
    if (results.get("dataset") or {}).get("synthetic"):
        raise HTTPException(status_code=409, detail="The card is disabled while test weather data is in use.")
    company = company_settings(session, settings)
    doc = AssessmentDoc.model_validate(a.doc)
    step = doc.card_next_step if next_step is None else next_step
    png = build_client_card(doc, results, company, next_step=step[:120], public_url=settings.estimate_url)
    return Response(png, media_type="image/png", headers=_download_name("roof-check", a, "png", inline=True))


# ---- bill of materials export: the generated list with the owner's edits, as the pricing results hold it.
# Round 3 (E-14): a header block (customer, project, date, system), the spec or model per line, the lines grouped
# by category in the customer sections' order, and the pack rounding noted where the item is sold by the roll or box.

BOM_COLUMNS = ["category", "code", "item", "spec / model", "supplier", "qty", "unit", "packs", "role", "note"]
BOM_CATEGORY_ORDER = ["Solar Panel", "Inverter", "Battery", "All-in-one System", "Mounting", "Wires and Terminations", "Protective Devices",
                      "Enclosures and Raceways", "Grounding", "Accessories", "Consumables"]
BOM_KIND_LABEL = {"net_metering": "net metering, no battery", "combination": "net metering with a battery", "off_grid": "no export, battery with the grid as backup"}
_PACK_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:m|meter)\b", re.I)


def _pack_note(qty: float, unit: str, sold_as: str, name: str) -> str:
    """"2 × 150 m box" when the item is sold by a pack whose size the list states (sold-as text or the name), else blank."""
    if (unit or "").lower() != "m":
        return ""
    text = f"{sold_as or ''} {name or ''}"
    m = _PACK_RE.search(sold_as or "") or _PACK_RE.search(name or "")
    if not m or "box" not in text.lower() and "roll" not in text.lower():
        return ""
    size = float(m.group(1))
    if size <= 0:
        return ""
    import math
    packs = int(math.ceil(qty / size - 1e-9))
    kind = "box" if "box" in text.lower() else "roll"
    return f"{packs} × {size:g} m {kind}"


def _bom_header(a: Assessment, results: dict) -> list[list]:
    """The block above the lines: who, which project, when, what system (as the pricing results hold it)."""
    pricing = results.get("pricing") or {}
    sizing = results.get("sizing") or {}
    tot = pricing.get("totals") or {}
    lines = pricing.get("lines") or []
    panel = next((l for l in lines if l.get("role") == "panel"), None)
    inverter = next((l for l in lines if l.get("role") == "inverter"), None)
    battery = next((l for l in lines if l.get("role") == "battery"), None)
    system = f"{float(tot.get('kwp') or 0):.2f} kWp, {BOM_KIND_LABEL.get(str(sizing.get('kind') or ''), 'solar PV system')}"
    if panel:
        system += f"; {int(float(panel.get('qty') or 0))} × {panel.get('name') or panel.get('code')}"
    if inverter:
        system += f"; inverter {int(float(inverter.get('qty') or 0))} × {inverter.get('name') or inverter.get('code')}"
    if battery:
        system += f"; battery {int(float(battery.get('qty') or 0))} × {battery.get('name') or battery.get('code')}"
    computed = str(results.get("computed_at") or "")[:10]
    return [
        ["Bill of materials"],
        ["Customer", a.customer_name or ""],
        ["Project", f"#{a.id} P-{a.created_at.year}-{a.id:04d}", a.address or ""],
        ["Date", computed or utcnow().date().isoformat(), "calculated; the quantities are the generated list with the owner's edits"],
        ["System", system],
        ["Packs", "the packs column rounds a length up to the roll or box the item is sold by; wastage is priced, not added to the quantity"],
        [],
    ]


def _bom_rows(a: Assessment, session: Session) -> tuple[list[list], list[list]]:
    """The header block and one row per BOM line (category, code, item, spec, supplier, qty, unit, packs, role, note),
    grouped by category; the one stale rule applies, and a record without priced results is refused the same way the
    proposal is. A role without an item (NO-ITEM-...) stays on the list with its quantity and the note that says so."""
    results = _fresh_results(a)
    pricing = results.get("pricing") or {}
    if not pricing.get("available"):
        raise HTTPException(status_code=409, detail="Calculate first. Pricing needs the panel linked to the materials list.")
    catalog = load_catalog(session, include_inactive=True)
    order = {c: i for i, c in enumerate(BOM_CATEGORY_ORDER)}
    rows: list[list] = []
    for l in pricing.get("lines") or []:
        qty = float(l.get("qty") or 0)
        code = str(l.get("code") or "")
        item = catalog.get(code)
        found = l.get("found", True)
        category = l.get("category") or (item.category if item else "") or ("(no item yet)" if code.startswith("NO-ITEM-") else "(not in the list)")
        name = l.get("name") if found else ("no item in the materials list for this role" if code.startswith("NO-ITEM-") else "not in the materials list")
        rows.append([
            category, code, name or "", (item.spec if item else "") or "", l.get("supplier") or "",
            int(qty) if qty.is_integer() else round(qty, 2), l.get("unit") or "", _pack_note(qty, l.get("unit") or "", item.sold_as if item else "", item.name if item else ""),
            l.get("role") or "", l.get("note") or "",
        ])
    rows.sort(key=lambda r: order.get(r[0], len(order)))
    return _bom_header(a, results), rows


@router.get("/{assessment_id}/bom.csv")
def bom_csv(assessment_id: int, session: Session = Depends(get_session)) -> Response:
    """The bill of materials as CSV (UTF-8 with a byte-order mark, so spreadsheets open it as typed): the header
    block, then the column header, then the lines grouped by category (a category row before each group)."""
    import csv
    import io

    a = _get(session, assessment_id)
    header, rows = _bom_rows(a, session)
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerows(header)
    w.writerow(BOM_COLUMNS)
    last = None
    for r in rows:
        if r[0] != last:
            w.writerow([f"— {r[0]} —"])
            last = r[0]
        w.writerow(r)
    return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv; charset=utf-8", headers=_download_name("bom", a, "csv"))


@router.get("/{assessment_id}/bom.xlsx")
def bom_xlsx(assessment_id: int, session: Session = Depends(get_session)) -> Response:
    """The bill of materials as a workbook: one sheet, the header block, a bold column header, a shaded row per
    category group, the columns sized to read."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    a = _get(session, assessment_id)
    header, rows = _bom_rows(a, session)
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    for h in header:
        ws.append(h)
    ws["A1"].font = Font(bold=True, size=13)
    for r in range(2, len(header)):
        ws.cell(row=r, column=1).font = Font(bold=True)
    ws.append(BOM_COLUMNS)
    head_row = ws.max_row
    for c in ws[head_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FAF4E1")
    last = None
    for r in rows:
        if r[0] != last:
            ws.append([r[0]])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="F2EFE6")
            last = r[0]
        ws.append(r)
    for i, width in enumerate([22, 20, 44, 30, 16, 8, 7, 16, 18, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=head_row + 1, min_col=6, max_col=6):
        for c in row:
            c.alignment = Alignment(horizontal="right")
    ws.freeze_panes = f"A{head_row + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=_download_name("bom", a, "xlsx"))
