"""Import the materials workbook (PLD_Materials_DB.xlsx) into catalog data and a PricingConfig.

The app is the master after import; re-importing updates prices by item code and
adds new codes without touching items that are not in the workbook.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import openpyxl

from .catalog import ELECTRICAL_FIELDS, Catalog, Item, Supplier, certifications_in_remarks, electrical_from_remarks, infer_grid_interactive, parse_panel_dims
from .config import CategoryRule, GroundTask, PricingConfig

# Optional electrical columns on MATERIALS DB (contract C4): matched by header text, in any column, when present.
# The bundled workbook has none; figures then come from the remarks (see catalog.electrical_from_remarks) or stay blank.
ELECTRICAL_HEADERS: dict[str, tuple[str, ...]] = {
    "grid_interactive": ("gridinteractive", "gridtie", "exportallowed"),
    "certifications": ("certifications", "certification", "certificate", "listing"),
    "max_pv_voltage_v": ("maxpvvoltage", "maxpvvoltagev", "pvvoltagemax", "maxdcvoltage", "maxinputvoltage"),
    "mppt_min_v": ("mpptmin", "mpptminv", "mpptrangemin", "mpptlow"),
    "mppt_max_v": ("mpptmax", "mpptmaxv", "mpptrangemax", "mppthigh"),
    "mppt_count": ("mpptcount", "mppts", "numberofmppt", "mpptinputs"),
    "mppt_max_a": ("mpptmaxa", "maxcurrentpermppt", "mpptcurrent", "currentpermppt"),
    "ac_input_a": ("acinputa", "acinputcurrent", "maxacinput", "maxacinputcurrent"),
    "battery_max_a": ("batterymaxa", "maxbatterycurrent", "batterycurrent", "batterychargecurrent"),
    "has_transfer_switch": ("transferswitch", "builtintransferswitch", "builtinats", "internalats", "hasats", "owntransferswitch"),
    "continuous_a": ("continuousa", "continuouscurrent", "continuousdischarge", "continuousdischargecurrent"),
    "voc_v": ("voc", "vocv", "opencircuitvoltage"),
    "vmp_v": ("vmp", "vmpv", "vmpp"),
    "isc_a": ("isc", "isca", "shortcircuitcurrent"),
    "imp_a": ("imp", "impa", "impp"),
    "temp_coeff_voc_pct": ("tempcoeffvoc", "tempcoeffvocpct", "voctempcoeff", "betavoc"),
    "temp_coeff_isc_pct": ("tempcoeffisc", "tempcoeffiscpct", "isctempcoeff", "alphaisc"),
    # round 12: the datasheet fields, so the materials workbook may carry the same columns the datasheet workbooks do
    "max_system_voltage_v": ("maxsystemvoltage", "maxsystemvoltagev", "maxdcsystemvoltage", "systemvoltage"),
    "inverter_type": ("invertertype", "typesofinverter", "typeofinverter"),
    "phase": ("phase", "phases"),
    "battery_class": ("batteryclass", "batteryvoltageclass", "lvhv"),
    "charge_v_max": ("chargevmax", "maxchargevoltage", "maxchargevoltagev", "chargevoltagemax"),
    "charge_a_max": ("chargeamax", "maxchargecurrent", "maxrecommendedchargecurrent", "chargecurrentmax"),
    "mppt_currents_a": ("mpptcurrents", "mpptcurrentsa", "mpptinputcurrents"),
    "battery_inputs": ("batteryinputs", "batteryports"),
    "nominal_v": ("nominalv", "nominalvoltage", "batteryvoltage"),
    "capacity_ah": ("capacityah", "capacity", "ah"),
    "discharge_a_recommended": ("dischargearecommended", "recommendeddischargecurrent", "recommendeddischarge"),
    # round 13: the design analysis's item figures, so a workbook may carry them as columns (brief 2.3)
    "overall_area_mm2": ("overallarea", "overallareamm2", "conductorarea", "insulatedarea"),
    "inner_diameter_mm": ("innerdiameter", "innerdiametermm", "insidediameter", "conduitinnerdiameter"),
    "insulation_c": ("insulation", "insulationc", "insulationrating", "insulationtemperature"),
    "ampacity_a": ("ampacity", "ampacitya", "ratedampacity", "cableampacity"),
    "fault_current_a": ("faultcurrent", "faultcurrenta", "maxoutputfaultcurrent", "faultcontribution"),
    "aic_ka": ("aic", "aicka", "interruptingrating", "interruptingcapacity"),
}
_INT_FIELDS = {"mppt_count", "phase", "battery_inputs"}
_TEXT_FIELDS = {"certifications", "inverter_type", "battery_class", "mppt_currents_a"}


def _norm_header(v: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(v or "").lower().split("(")[0])


def electrical_columns(header_row: tuple) -> dict[str, int]:
    """Column index per electrical field, from a header row; fields whose header is absent are left out."""
    found: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        key = _norm_header(cell)
        if not key:
            continue
        for field_name, aliases in ELECTRICAL_HEADERS.items():
            if key in aliases and field_name not in found:
                found[field_name] = idx
    return found


def _bool(v: Any) -> Optional[bool]:
    s = _s(v).lower()
    if s in ("yes", "y", "true", "1", "grid-tie", "grid tie"):
        return True
    if s in ("no", "n", "false", "0", "off-grid", "off grid"):
        return False
    return None


def electrical_values(row: tuple, cols: dict[str, int], category: str, name: str, spec: str, remarks: str) -> dict:
    """The electrical fields for one row: the workbook column when present and filled, else what the remarks say,
    else None. `grid_interactive` is inferred from the name and remarks when no column answers it."""
    out: dict = dict.fromkeys(ELECTRICAL_FIELDS, None)
    for k in _TEXT_FIELDS:
        out[k] = ""
    out.update(electrical_from_remarks(category, name, spec, remarks))
    if category == "Inverter":
        out["certifications"] = certifications_in_remarks(remarks)
        out["grid_interactive"] = infer_grid_interactive(name, remarks)
    for field_name, idx in cols.items():
        v = row[idx] if idx < len(row) else None
        if v is None or v == "":
            continue
        if field_name in ("grid_interactive", "has_transfer_switch"):
            b = _bool(v)
            if b is not None:
                out[field_name] = b
        elif field_name in _TEXT_FIELDS:
            out[field_name] = _s(v)
        elif field_name in _INT_FIELDS:
            out[field_name] = int(_f(v))
        else:
            out[field_name] = _f(v)
    return out


@dataclass
class ImportResult:
    catalog: Catalog
    config: PricingConfig
    item_count: int
    supplier_count: int
    warnings: list[str]


def _f(v: Any, default: float = 0.0) -> float:
    if v is None or v == "":
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _s(v: Any) -> str:
    return "" if v is None else str(v).strip()


def read_workbook(path: str | Path) -> ImportResult:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    warnings: list[str] = []
    cfg = PricingConfig()

    # ---- suppliers
    suppliers: dict[str, Supplier] = {}
    if "SUPPLIERS" in wb.sheetnames:
        ws = wb["SUPPLIERS"]
        for row in ws.iter_rows(min_row=3, values_only=True):
            name = _s(row[0])
            if not name:
                continue
            suppliers[name] = Supplier(
                name=name, pickup_address=_s(row[1]), dealer_discount=_f(row[3]), payment_fee=_f(row[4]),
                delivers_free=_s(row[5]).lower() == "yes", price_list_date=_s(row[6]), prices_note=_s(row[7]),
                warranty=_s(row[8]), remarks=_s(row[9]),
            )

    # ---- items
    items: dict[str, Item] = {}
    ws = wb["MATERIALS DB"]
    header = next(ws.iter_rows(min_row=3, max_row=3, values_only=True), ())
    elec_cols = electrical_columns(header)
    for row in ws.iter_rows(min_row=4, values_only=True):
        code = _s(row[0])
        if not code or not re.match(r"^[A-Z]+-[A-Z]+-\d+$", code):
            continue
        spec = _s(row[4])
        length, width = parse_panel_dims(spec)
        rating = row[10]
        category, name = _s(row[1]), _s(row[3])
        remarks = _s(row[29]) if len(row) > 29 else ""
        items[code] = Item(
            code=code, category=category, supplier=_s(row[2]), name=name, spec=spec,
            unit=_s(row[5]) or "pc", sold_as=_s(row[6]) or "pc", list_price=_f(row[7]),
            rating=None if rating in (None, "") else _f(rating), rating_unit=_s(row[11]),
            weight_kg=_f(row[13]), volume_m3=_f(row[14]), weight_source=_s(row[15]), storage=_f(row[21]),
            price_list_date=_s(row[27]) if len(row) > 27 else "", remarks=remarks,
            panel_length_m=length if category == "Solar Panel" else None,
            panel_width_m=width if category == "Solar Panel" else None,
            **electrical_values(row, elec_cols, category, name, spec, remarks),
        )
    for it in items.values():
        if it.category == "Solar Panel" and it.panel_length_m is None:
            warnings.append(f"{it.code}: no panel dimensions in the spec; fill length and width in the materials page.")
        if it.supplier and it.supplier not in suppliers:
            warnings.append(f"{it.code}: supplier {it.supplier!r} not on the SUPPLIERS sheet.")

    # ---- DRIVERS: truck, handling, job-level, categories
    if "DRIVERS" in wb.sheetnames:
        ws = wb["DRIVERS"]
        vals = {_s(r[0]): r[1] for r in ws.iter_rows(min_row=1, max_row=30, values_only=True) if r and r[0] is not None}
        t = cfg.truck
        t.payload_kg = _f(vals.get("Payload"), t.payload_kg)
        t.cargo_volume_m3 = _f(vals.get("Cargo volume"), t.cargo_volume_m3)
        t.ownership_per_trip_day = _f(vals.get("Ownership cost per trip-day"), t.ownership_per_trip_day)
        t.driver_per_trip_day = _f(vals.get("Driver per trip-day"), t.driver_per_trip_day)
        t.helper_per_trip_day = _f(vals.get("Helper per trip-day"), t.helper_per_trip_day)
        t.diesel_price = _f(vals.get("Diesel price"), t.diesel_price)
        t.fuel_economy_km_per_l = _f(vals.get("Fuel economy, loaded"), t.fuel_economy_km_per_l)
        t.maintenance_per_km = _f(vals.get("Maintenance (tires, PMS) per km"), t.maintenance_per_km)
        t.trip_days_per_run = _f(vals.get("Trip-days for the run"), t.trip_days_per_run)
        h = cfg.handling
        h.helper_day_rate = _f(vals.get("Helper day rate (incl. meal)"), h.helper_day_rate)
        h.typical_job_helper_hours = _f(vals.get("Handling hours, typical 6 kW job"), h.typical_job_helper_hours)
        h.typical_fill = _f(vals.get("Typical fill of the truck (from REF BASKET)"), h.typical_fill)
        j = cfg.job
        j.agent_commission = _f(vals.get("Agent commission (of direct cost)"), j.agent_commission)
        j.vat = _f(vals.get("VAT"), j.vat)
        j.freight_markup = _f(vals.get("Markup on freight (job line)"), j.freight_markup)
        cats: list[CategoryRule] = []
        started = False
        for r in ws.iter_rows(min_row=1, values_only=True):
            if _s(r[0]) == "Category" and _s(r[1]) == "Markup tier":
                started = True
                continue
            if started and _s(r[0]):
                cats.append(CategoryRule(name=_s(r[0]), markup_tier=_f(r[1]), wastage=_f(r[2])))
        if cats:
            cfg.categories = cats

    # ---- ROUTE: stops and matrices (km B13:G18, toll B22:G27)
    if "ROUTE" in wb.sheetnames:
        ws = wb["ROUTE"]
        stops = [_s(ws.cell(row=r, column=2).value) for r in range(4, 10)]
        stops = [s for s in stops if s]
        km = [[_f(ws.cell(row=r, column=c).value) for c in range(2, 2 + len(stops))] for r in range(13, 13 + len(stops))]
        toll = [[_f(ws.cell(row=r, column=c).value) for c in range(2, 2 + len(stops))] for r in range(22, 22 + len(stops))]
        if stops and len(km) == len(stops):
            cfg.route.stops, cfg.route.km, cfg.route.toll = stops, km, toll
            cfg.route.reference_site_km = km[0][-1]

    # ---- LABOR RATES
    if "LABOR RATES" in wb.sheetnames:
        ws = wb["LABOR RATES"]
        vals = {_s(r[0]): r[1] for r in ws.iter_rows(min_row=1, max_row=75, values_only=True) if r and r[0] is not None}
        lr = cfg.labor
        lr.team_lead_day = _f(vals.get("Team lead (electrician)"), lr.team_lead_day)
        lr.skilled_day = _f(vals.get("Skilled installer"), lr.skilled_day)
        lr.laborer_day = _f(vals.get("Laborer"), lr.laborer_day)
        lr.allowance_included = _f(vals.get("Allowance included in each rate above"), lr.allowance_included)
        lr.owner_day = _f(vals.get("You, if on site"), lr.owner_day)
        lr.paid_hours = _f(vals.get("Paid hours per day"), lr.paid_hours)
        lr.nonproductive_hours = _f(vals.get("Non-productive hours (toolbox talk, tool setup, end-of-day cleanup)"), lr.nonproductive_hours)
        lr.pay_unit_days = _f(vals.get("Pay unit (1 = whole days, 0.5 = half days)"), lr.pay_unit_days)
        rf = cfg.roof
        rf.setup_mh = _f(vals.get("Roof setup per job (part of the mounting hours)"), rf.setup_mh)
        rf.mounting_mh_per_panel = _f(vals.get("Mounting per panel (L-feet, rails, clamps, panel)"), rf.mounting_mh_per_panel)
        rf.wiring_mh_per_panel = _f(vals.get("Wiring per panel (panel-to-panel, string runs on the roof)"), rf.wiring_mh_per_panel)
        rf.simple_roof_factor = _f(vals.get("Roof factor: simple roof"), rf.simple_roof_factor)
        g = cfg.ground
        g.mounting_mh_per_unit = _f(vals.get("Mounting: man-hours per weighted unit"), g.mounting_mh_per_unit)
        g.wiring_mh_per_unit = _f(vals.get("Wiring: man-hours per weighted unit"), g.wiring_mh_per_unit)
        g.hybrid_pace = _f(vals.get("Hybrid ground pace"), g.hybrid_pace)
        cfg.hauling.max_kg_per_person = _f(vals.get("Max weight carried per person"), cfg.hauling.max_kg_per_person)
        # ground task weights, rows 39-50: label, mounting weight, wiring weight, mh, unit
        keys = [t.key for t in g.tasks]
        new_tasks: list[GroundTask] = []
        for r in ws.iter_rows(min_row=39, max_row=50, values_only=True):
            if not r or not _s(r[0]):
                continue
            idx = len(new_tasks)
            if idx < len(keys):
                new_tasks.append(GroundTask(key=keys[idx], label=_s(r[0]), mounting_weight=_f(r[1]), wiring_weight=_f(r[2]), unit=_s(r[4]) or g.tasks[idx].unit))
        if len(new_tasks) == len(keys):
            g.tasks = new_tasks

    # ---- MOB-DEMOB and TOOLS
    if "MOB-DEMOB" in wb.sheetnames:
        ws = wb["MOB-DEMOB"]
        vals = {_s(r[0]): r[1] for r in ws.iter_rows(min_row=1, max_row=30, values_only=True) if r and r[0] is not None}
        m = cfg.mobdemob
        m.vehicle_ownership_per_day = _f(vals.get("Ownership cost per day used"), m.vehicle_ownership_per_day)
        m.running_cost_per_km = _f(vals.get("Running cost per km"), m.running_cost_per_km)
        m.base_to_site_km = _f(vals.get("One-way km, base to site"), m.base_to_site_km)
        m.toll_per_round_trip = _f(vals.get("Toll per round trip"), m.toll_per_round_trip)
        m.one_way_travel_hours = _f(vals.get("One-way travel time"), m.one_way_travel_hours)
        m.packaging_disposal = _f(vals.get("Packaging and debris haul-out and disposal"), m.packaging_disposal)
    if "TOOLS" in wb.sheetnames:
        ws = wb["TOOLS"]
        for r in ws.iter_rows(min_row=1, values_only=True):
            if r and _s(r[0]) == "Tool charge per installation day":
                cfg.tools.charge_per_installation_day = _f(r[1], cfg.tools.charge_per_installation_day)

    # ---- JOB sheet defaults and job-level costs
    if "JOB" in wb.sheetnames:
        ws = wb["JOB"]
        vals = {_s(r[0]): r[1] for r in ws.iter_rows(min_row=1, max_row=140, values_only=True) if r and r[0] is not None}
        d = cfg.job_defaults
        d.roof_factor = _f(vals.get("Roof difficulty factor"), d.roof_factor)
        d.roof_closed_days = int(_f(vals.get("Roof closed within (days)"), d.roof_closed_days))
        d.max_days = int(_f(vals.get("Max installation days"), d.max_days))
        d.max_pairs = int(_f(vals.get("Max pairs"), d.max_pairs))
        d.battery_haul_hours = _f(vals.get("Battery hauling hours per pack"), d.battery_haul_hours)
        j = cfg.job
        j.services_markup = _f(vals.get("Markup on labor, mob/demob, tools, PPE and the seal (input)"), j.services_markup)
        j.ppe_per_person_day = _f(vals.get("PPE per person-day (input)"), j.ppe_per_person_day)
        j.ocm_share = _f(vals.get("OCM share of markup (input)"), j.ocm_share)
        j.pee_seal = _f(vals.get("PEE sign and seal"), j.pee_seal)
        j.lgu_permit_cfei = _f(vals.get("LGU electrical permit and CFEI"), j.lgu_permit_cfei)
        erc = _f(vals.get("ERC Certificate of Compliance fee (net metering)"))
        meter = _f(vals.get("Bi-directional meter difference (net metering)"))
        if erc > 0:
            j.erc_coc_fee = erc
        if meter > 0:
            j.bidirectional_meter_fee = meter

    cfg.imported_from = path.name
    cfg.imported_at = datetime.now(timezone.utc).isoformat()
    resolve_roles(cfg, items, warnings)
    return ImportResult(Catalog(items, suppliers), cfg, len(items), len(suppliers), warnings)


def resolve_roles(cfg: PricingConfig, items: dict[str, Item], warnings: list[str]) -> None:
    """Fill gauge maps from item names when the default codes are missing from the DB."""
    def find(pattern: str, prefer_unit: Optional[str] = None) -> Optional[str]:
        rx = re.compile(pattern, re.I)
        hits = [i for i in items.values() if rx.search(i.name)]
        if prefer_unit:
            hits = sorted(hits, key=lambda i: (i.unit != prefer_unit, i.list_price))
        else:
            hits = sorted(hits, key=lambda i: i.list_price)
        return hits[0].code if hits else None

    r = cfg.roles
    for gauge in list(r.thhn):
        if r.thhn[gauge] not in items:
            code = find(rf"^THHN {re.escape(gauge)}mm2 EURO WIRES$", "m") or find(rf"^THHN {re.escape(gauge)}mm2$", "m")
            if code:
                r.thhn[gauge] = code
            else:
                warnings.append(f"No THHN {gauge} mm2 item found; role left as {r.thhn[gauge]}.")
    for gauge in list(r.battery_cable_pair):
        if r.battery_cable_pair[gauge] not in items:
            code = find(rf"^{gauge}mm2 - 1M w/ LUG PAIR$")
            if code:
                r.battery_cable_pair[gauge] = code
    for role in ("rail", "l_foot", "end_clamp", "mid_clamp", "splice", "mc4_pair", "dc_breaker", "dc_spd", "battery_breaker_fallback", "ats", "ac_breaker", "ac_spd", "enclosure", "cable_tray", "conduit", "ground_rod", "earth_lug", "sealant",
                 "ac_disconnect", "array_bonding_wire"):
        if getattr(r, role) not in items:
            warnings.append(f"Role {role}: default code {getattr(r, role)} is not in the DB; set it on the materials page.")
    # the grid default: the first grid-interactive hybrid in the catalogue (code order), else blank = the cheapest that fits
    excluded = [w.lower() for w in r.inverter_exclude_words]
    grid_ok = sorted(
        (i for i in items.values() if i.category == "Inverter" and i.is_hybrid_inverter and i.grid_interactive is True
         and not any(w in f"{i.name} {i.spec}".lower() for w in excluded)),
        key=lambda i: i.code,
    )
    # the owner's default (the eco-hybrid) stays the grid default when it may export; else the first grid-interactive hybrid
    preferred = next((i for i in items.values() if i.code in (r.default_inverter_code_grid, r.default_inverter_code_offgrid) and i.grid_interactive is True), None)
    r.default_inverter_code_grid = preferred.code if preferred else (grid_ok[0].code if grid_ok else "")
    if not grid_ok:
        warnings.append("No grid-interactive hybrid inverter in the workbook; net-metering jobs take the cheapest grid-interactive unit that fits once one is marked on the Materials page.")
    if r.default_inverter_code_offgrid and r.default_inverter_code_offgrid not in items:
        warnings.append(f"Default off-grid inverter {r.default_inverter_code_offgrid} is not in the DB; set it under BOM item roles.")
