"""The engineering drawings: results["geometry"] behind the plan of each face, the plan and Gantt flowables, the
drawings in the roof check, the proposal and the program of works, and the bill of materials export."""
import io
import shutil
import subprocess
from copy import deepcopy
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from reportlab.graphics import renderPDF
from reportlab.lib.units import mm

from solarapp.config import Settings
from solarapp.data_download.cli import write_synthetic
from solarapp.main import create_app
from solarapp.reports.customer_pdf import build_customer_pdf
from solarapp.reports.drawings import gantt_drawing, plan_drawing
from solarapp.schemas import AssessmentDoc

STALE = "Inputs changed since the last calculation. Calculate again first."


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    root = tmp_path_factory.mktemp("data")
    write_synthetic(root, (14.0, 14.5, 121.0, 121.5), 0.25)  # Laguna
    settings = Settings(data_dir=root, app_username="u", app_password="p", secret_key="s" * 32, cookie_secure=False, static_dir=root / "nostatic")
    app = create_app(settings)
    with TestClient(app) as c:
        c.post("/api/auth/login", json={"username": "u", "password": "p"})
        yield c


# a south rectangle with a mango tree to the west-southwest and an east hip face beside a firewall on its left
PILA_DOC = {
    "customer_name": "Maria Santos", "address": "Brgy. Labuin, Pila, Laguna", "lat": 14.2335, "lon": 121.3645,
    "faces": [
        {"id": "f1", "name": "Main roof (south)", "shape": "rect", "length_m": 9.0, "width_m": 5.0, "tilt_deg": 18, "azimuth_deg": 180,
         "obstacles": [{"id": "t1", "label": "Mango tree", "direction_deg": 250, "elevation_deg": 24, "width_deg": 50, "share": 0.5}]},
        {"id": "f2", "name": "Kitchen roof (east)", "shape": "hip", "length_m": 7.0, "width_m": 4.0, "ridge_m": 3.0, "tilt_deg": 15, "azimuth_deg": 90,
         "walls": [{"id": "w1", "edge": "left", "height_m": 1.5, "gap_m": 0}]},
    ],
    "panels": [{"id": "p1", "name": "Blue Carbon 585W", "watt_peak": 585, "length_m": 2.278, "width_m": 1.134, "code": "BC-PNL-001"}],
    "reading_sets": [
        {"id": "s1", "face_id": "f1", "measured_at": "2026-10-03T11:30:00", "ambient_temp_c": 33, "sky_condition": "clear",
         "readings": [{"irradiance_wm2": 903, "power_w": 36.1, "module_temp_c": 57}, {"irradiance_wm2": 910, "power_w": 36.4, "module_temp_c": 58}, {"irradiance_wm2": 896, "power_w": 35.8, "module_temp_c": 57}]},
    ],
    "audit": {
        "appliances": [
            {"id": "ref", "name": "Refrigerator", "category": "refrigerator", "input_power_w": 150, "quantity": 1, "windows": [{"start": "06:00", "end": "06:00"}]},
            {"id": "ac", "name": "Split inverter AC", "category": "aircon_inverter", "input_power_w": 1200, "quantity": 1, "windows": [{"start": "20:00", "end": "06:00"}]},
            {"id": "led", "name": "LED bulb", "category": "lighting", "input_power_w": 9, "quantity": 8, "windows": [{"start": "18:00", "end": "06:00"}]},
            {"id": "tv", "name": "TV", "category": "other", "input_power_w": 90, "quantity": 1, "windows": [{"start": "18:00", "end": "23:00"}]},
        ],
        "bills": [{"id": "b1", "billing_month": "2026-09", "kwh": 310, "days": 30, "amount_php": 3980.5, "utility": "FLECO"}],
        "reconcile": True,
        "system": {"kind": "combination"},
    },
    "program": {"stage": "quoted", "signing_date": "2026-10-12"},
}


def _computed(client: TestClient, doc: dict) -> dict:
    aid = client.post("/api/assessments", json=doc).json()["id"]
    r = client.post(f"/api/assessments/{aid}/compute")
    assert r.status_code == 200, r.text
    return r.json()


def _pdf_text(pdf: bytes) -> str:
    if not shutil.which("pdftotext"):
        pytest.skip("pdftotext is not installed")
    return subprocess.run(["pdftotext", "-layout", "-", "-"], input=pdf, capture_output=True, check=True).stdout.decode()


def _pages(text: str) -> int:
    return text.count("\f")


def test_geometry_follows_the_fitted_layout(client):
    res = _computed(client, PILA_DOC)["results"]
    geo = res["geometry"]
    assert [g["face_id"] for g in geo] == ["f1", "f2"] and [g["name"] for g in geo] == ["Main roof (south)", "Kitchen roof (east)"]
    sel = next(p for p in res["panels"] if p["panel"]["id"] == res["selected_panel_id"])
    for g in geo:
        lr = sel["faces"][g["face_id"]]
        assert len(g["panels"]) == lr["count"] == g["count"] and g["orientation"] == lr["best"]["orientation"] and g["gross"] == lr["gross"]
        rows: dict[int, int] = {}
        for p in g["panels"]:
            rows[p["row"]] = rows.get(p["row"], 0) + 1
        assert [rows[k] for k in sorted(rows)] == [n for n in lr["best"]["rows"] if n > 0]
        for k in ("eave", "ridge", "left", "right"):
            assert abs(g["cuts"][k] - lr["cuts"][k]) < 1e-3
        assert g["panel_length_m"] == 2.278 and g["panel_width_m"] == 1.134 and g["setback_m"] == 0.6
    main, kitchen = geo
    assert main["shape"] == "rect" and main["compass"] == "south" and len(main["outline"]) == 4
    tree = next(o for o in main["obstacles"] if o["kind"] == "shade")
    assert tree["edge"] == "left" and tree["x"] == 0.0 and tree["label"].startswith("Mango tree, west-southwest, 24° up")
    assert kitchen["shape"] == "hip" and kitchen["outline"] == [[0, 0], [7, 0], [5, 4], [2, 4]]
    wall = next(o for o in kitchen["obstacles"] if o["kind"] == "wall")
    assert wall["edge"] == "left" and abs(wall["w"] - res["shade"]["f2"]["walls"][0]["strip_m"]) < 1e-3 and wall["h"] == 4.0
    assert all(p["x"] >= wall["w"] - 1e-6 for p in kitchen["panels"])
    # the sized system: faces in order, rows from the eave up, strings by the BOQ's rule
    sized = res["sizing"]["panels"]
    assert sum(g["used"] for g in geo) == sized and main["used"] == min(sized, main["count"])
    used = [p for g in geo for p in g["panels"] if p["used"]]
    assert len(used) == sized and all("string" in p for p in used)
    assert all("string" not in p for g in geo for p in g["panels"] if not p["used"])
    choices = res["pricing"]["choices"]
    assert max(p["string"] for p in used) == choices["strings"]
    assert max(sum(1 for p in used if p["string"] == s) for s in range(1, choices["strings"] + 1)) == choices["panels_per_string"]
    # the layout results next to it are what they were
    assert sel["faces"]["f1"]["best"]["rows"] == [3, 3, 3] and sel["total_count"] == main["count"] + kitchen["count"]


def test_geometry_without_an_audit_marks_nothing(client):
    doc = deepcopy(PILA_DOC)
    doc["audit"] = {"appliances": [], "bills": [], "reconcile": True, "system": {"kind": "combination"}}
    res = _computed(client, doc)["results"]
    assert res["sizing"] is None
    assert all(g["used"] is None for g in res["geometry"]) and all("used" not in p for g in res["geometry"] for p in g["panels"])


def test_plan_and_gantt_flowables_render():
    face = {
        "face_id": "f", "name": "Main roof", "shape": "hip", "eave_m": 10, "slope_m": 4.5, "ridge_m": 4, "azimuth_deg": 180, "tilt_deg": 18, "compass": "south",
        "orientation": "landscape", "outline": [[0, 0], [10, 0], [7, 4.5], [3, 4.5]], "usable": [[0.561, 0.3], [9.439, 0.3], [6.839, 4.2], [3.161, 4.2]],
        "panels": [{"x": 1.563, "y": 0.3, "w": 2.278, "h": 1.134, "row": 1, "n": 1, "used": True}, {"x": 3.861, "y": 0.3, "w": 2.278, "h": 1.134, "row": 1, "n": 2, "used": False}],
        "count": 2, "gross": 2, "left_out": 0, "used": 1, "cuts": {"eave": 0, "ridge": 0, "left": 1.2, "right": 0},
        "obstacles": [{"kind": "wall", "edge": "left", "x": 0, "y": 0, "w": 1.2, "h": 4.5, "label": "Wall on the left side, 1 m above the roof: no panels within 1.2 m of it"},
                      {"kind": "shade", "edge": "eave", "x": 5, "y": 0, "w": 0, "h": 0, "cls": "small", "label": "Mango tree, south, 24° up: small loss"}],
    }
    for highlight in (False, True):
        d = plan_drawing(face, 86, highlight_used=highlight, max_height_mm=40)
        assert abs(d.width - 86 * mm) < 1e-6 and 0 < d.height <= (40 + 30) * mm
        assert renderPDF.drawToString(d).startswith(b"%PDF")
    events = [
        {"key": "signing", "label": "Contract signing and downpayment", "date": "2026-10-12", "end": None, "kind": "milestone"},
        {"key": "permit_prep", "label": "Electrical plans, signed and sealed by the PEE", "date": "2026-10-12", "end": "2026-10-14", "kind": "task"},
        {"key": "down", "label": "Downpayment on signing", "date": "2026-10-12", "end": None, "kind": "payment_in", "amount": 136100},
        {"key": "install", "label": "Installation (1 day, crew of 4)", "date": "2026-10-22", "end": "2026-10-22", "kind": "task"},
        {"key": "meter", "label": "Electric company inspection; net metering meter installed, a long label that needs trimming to fit the column", "date": "2026-11-27", "end": None, "kind": "milestone"},
    ]
    g = gantt_drawing(events, 269, today=date(2026, 10, 20))
    assert abs(g.width - 269 * mm) < 1e-6 and g.height > 5 * 4.6 * mm
    assert renderPDF.drawToString(g).startswith(b"%PDF")
    assert renderPDF.drawToString(gantt_drawing([], 180)).startswith(b"%PDF")


def test_documents_carry_the_drawings(client):
    out = _computed(client, PILA_DOC)
    aid, res = out["id"], out["results"]
    from tests.conftest import real_weather
    real_weather(aid)
    q = client.get(f"/api/assessments/{aid}/quotation.pdf")
    assert q.status_code == 200
    text = _pdf_text(q.content)
    assert "YOUR ROOF, AS THE PANELS WILL SIT" in text and "Main roof (south)" in text and "Eave (lower edge)" in text
    assert "room for" in text and "Hatched: wall on the left side" in text
    assert _pages(text) <= 3   # the reference record prints on three pages
    p = client.get(f"/api/assessments/{aid}/program.pdf")
    assert p.status_code == 200
    ptext = _pdf_text(p.content)
    assert "Weeks from" in ptext and "milestone" in ptext and "payment from the customer" in ptext
    assert ptext.index("Weeks from") < ptext.index("Customer sees")   # the Gantt sits above the table
    # the roof check is off under test weather data; build it directly
    pdf = build_customer_pdf(AssessmentDoc.model_validate(out["doc"]), res, {"company_name": "Test Solar"})
    rtext = _pdf_text(pdf)
    assert "Your roof, as the panels will sit" in rtext and "Kitchen roof (east)" in rtext and "Measured on your roof" in rtext
    assert rtext.index("Your roof, as the panels will sit") < rtext.index("Measured on your roof")
    assert _pages(rtext) <= 2   # the reference record prints on two pages
    # a face with no panels gets one line
    doc = deepcopy(PILA_DOC)
    doc["faces"][1] = {**doc["faces"][1], "shape": "rect", "width_m": 0.8}
    out2 = _computed(client, doc)
    text2 = _pdf_text(build_customer_pdf(AssessmentDoc.model_validate(out2["doc"]), out2["results"], {"company_name": "Test Solar"}))
    assert "Kitchen roof (east) (7 × 0.8 m): no panels on this face." in text2


def test_bom_export_both_ways(client):
    out = _computed(client, PILA_DOC)
    aid, lines = out["id"], out["results"]["pricing"]["lines"]
    r = client.get(f"/api/assessments/{aid}/bom.csv")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"].startswith('attachment; filename="bom-Maria_Santos.csv"')
    rows = r.content.decode("utf-8-sig").splitlines()
    assert rows[0] == "code,item,supplier,qty,unit,role,note" and len(rows) == len(lines) + 1
    panel = next(l for l in rows[1:] if l.startswith("BC-PNL-001,"))
    assert ",Blue Carbon," in panel and ",pc,panel," in panel
    x = client.get(f"/api/assessments/{aid}/bom.xlsx")
    assert x.status_code == 200 and x.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert x.headers["content-disposition"].startswith('attachment; filename="bom-Maria_Santos.xlsx"')
    ws = load_workbook(io.BytesIO(x.content))["BOM"]
    assert [c.value for c in ws[1]] == ["code", "item", "supplier", "qty", "unit", "role", "note"] and ws.max_row == len(lines) + 1
    assert ws.freeze_panes == "A2" and ws["A1"].font.bold
    codes = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=4).value for i in range(2, ws.max_row + 1)}
    assert codes["BC-PNL-001"] == next(l["qty"] for l in lines if l["code"] == "BC-PNL-001")
    # the owner's edits travel with the export
    doc = deepcopy(out["doc"])
    doc["pricing"]["bom_edits"] = [{"code": "BC-PNL-001", "qty": 5, "note": ""}]
    r2 = client.post(f"/api/assessments/{aid}/compute", json=doc)
    assert r2.status_code == 200
    rows = client.get(f"/api/assessments/{aid}/bom.csv").content.decode("utf-8-sig").splitlines()
    assert any(l.startswith("BC-PNL-001,") and ",5,pc,panel," in l and "(edited)" in l for l in rows)
    # the one stale rule, and a record without pricing
    doc["notes"] = "changed"
    assert client.put(f"/api/assessments/{aid}", json=doc).json()["results_stale"] is True
    assert client.get(f"/api/assessments/{aid}/bom.csv").json()["detail"] == STALE
    assert client.get(f"/api/assessments/{aid}/bom.xlsx").json()["detail"] == STALE
    fresh = client.post("/api/assessments", json=PILA_DOC).json()["id"]
    assert client.get(f"/api/assessments/{fresh}/bom.xlsx").status_code == 409
    unpriced = deepcopy(PILA_DOC)
    unpriced["audit"] = {"appliances": [], "bills": [], "reconcile": True, "system": {"kind": "combination"}}
    u = _computed(client, unpriced)["id"]
    assert client.get(f"/api/assessments/{u}/bom.csv").status_code == 409
